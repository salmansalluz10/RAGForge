import re
from pathlib import Path
from typing import List, Optional
from dataclasses import dataclass
import pypdf
import docx
from fastapi import HTTPException, status


@dataclass
class ExtractedContent:
    text: str
    page_number: Optional[int] = None
    metadata: Optional[dict] = None


def normalize_text(text: str) -> str:
    """
    Clean and normalize extracted text:
    - Replace null characters and control sequences
    - Standardize whitespace and paragraph spacing
    - Strip trailing whitespaces
    """
    if not text:
        return ""
    # Strip null characters
    text = text.replace("\x00", "")
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Replace sequences of 3+ newlines with 2 newlines (preserve paragraph separation)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Replace multiple spaces/tabs with single space (except newlines)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


class BaseExtractor:
    def extract(self, file_path: Path) -> List[ExtractedContent]:
        raise NotImplementedError


class PDFExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> List[ExtractedContent]:
        results: List[ExtractedContent] = []
        try:
            reader = pypdf.PdfReader(str(file_path))
            total_pages = len(reader.pages)
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                cleaned = normalize_text(page_text)
                if cleaned:
                    results.append(
                        ExtractedContent(
                            text=cleaned,
                            page_number=page_idx + 1,
                            metadata={"total_pages": total_pages},
                        )
                    )
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")

        return results


class DocxExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> List[ExtractedContent]:
        results: List[ExtractedContent] = []
        try:
            doc = docx.Document(str(file_path))
            paragraphs = []

            # Extract body paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    paragraphs.append(p_text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        paragraphs.append(" | ".join(row_data))

            full_text = "\n\n".join(paragraphs)
            cleaned = normalize_text(full_text)
            if cleaned:
                results.append(
                    ExtractedContent(
                        text=cleaned,
                        page_number=1,
                        metadata={"paragraphs_count": len(paragraphs)},
                    )
                )
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

        return results


class TxtExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> List[ExtractedContent]:
        raw_bytes = file_path.read_bytes()
        # Attempt multiple standard encodings
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
        decoded_text = None

        for enc in encodings:
            try:
                decoded_text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if decoded_text is None:
            decoded_text = raw_bytes.decode("utf-8", errors="replace")

        cleaned = normalize_text(decoded_text)
        if not cleaned:
            return []

        return [ExtractedContent(text=cleaned, page_number=1, metadata={"encoding": enc})]


class ExtractorFactory:
    @staticmethod
    def get_extractor(file_type: str) -> BaseExtractor:
        normalized = file_type.lower().strip().lstrip(".")
        if normalized == "pdf":
            return PDFExtractor()
        elif normalized == "docx":
            return DocxExtractor()
        elif normalized == "txt":
            return TxtExtractor()
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"No extractor available for file type '{file_type}'.",
            )
