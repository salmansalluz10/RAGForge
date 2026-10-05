import io
import pytest
import docx
from pypdf import PdfWriter
from fastapi.testclient import TestClient
from app.main import app
from app.core.database import Base, engine

client = TestClient(app)


@pytest.fixture(autouse=True, scope="module")
def setup_database():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def get_auth_token(email: str = "processor@example.com") -> str:
    res = client.post("/api/v1/auth/register", json={
        "name": "Processor User",
        "email": email,
        "password": "Password123!",
    })
    return res.json()["access_token"]


def test_process_txt_document():
    token = get_auth_token("txt_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    sample_text = (
        "RAGForge is an enterprise AI document intelligence platform.\n\n"
        "It provides retrieval-augmented generation powered by pgvector and large language models.\n\n"
        "Documents are chunked into semantic segments with preserved metadata."
    )
    files = {"file": ("enterprise_overview.txt", io.BytesIO(sample_text.encode("utf-8")), "text/plain")}
    res_upload = client.post("/api/v1/documents/upload", headers=headers, files=files)
    assert res_upload.status_code == 201
    doc_id = res_upload.json()["id"]

    # Trigger processing
    res_process = client.post(f"/api/v1/documents/{doc_id}/process", headers=headers)
    assert res_process.status_code == 200
    data = res_process.json()
    assert data["processing_status"] == "COMPLETED"
    assert data["chunk_count"] >= 1
    assert data["error_message"] is None

    # Fetch chunks
    res_chunks = client.get(f"/api/v1/documents/{doc_id}/chunks", headers=headers)
    assert res_chunks.status_code == 200
    chunk_data = res_chunks.json()
    assert chunk_data["total"] >= 1
    first_chunk = chunk_data["chunks"][0]
    assert "RAGForge" in first_chunk["content"]
    assert first_chunk["chunk_metadata"]["document_id"] == doc_id
    assert first_chunk["chunk_metadata"]["file_type"] == "txt"


def test_process_docx_document():
    token = get_auth_token("docx_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Generate a real valid DOCX in memory
    doc = docx.Document()
    doc.add_heading("Financial Quarterly Analysis", level=1)
    doc.add_paragraph("Total revenue grew by 24% year-over-year to $45 million.")
    doc.add_paragraph("Operating margins expanded significantly due to cloud efficiencies.")

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    files = {
        "file": (
            "quarterly_report.docx",
            docx_buffer,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    res_upload = client.post("/api/v1/documents/upload", headers=headers, files=files)
    assert res_upload.status_code == 201
    doc_id = res_upload.json()["id"]

    # Trigger processing
    res_process = client.post(f"/api/v1/documents/{doc_id}/process", headers=headers)
    assert res_process.status_code == 200
    data = res_process.json()
    assert data["processing_status"] == "COMPLETED"
    assert data["chunk_count"] >= 1

    # Verify chunks
    res_chunks = client.get(f"/api/v1/documents/{doc_id}/chunks", headers=headers)
    assert res_chunks.status_code == 200
    chunks = res_chunks.json()["chunks"]
    assert any("revenue grew by 24%" in c["content"] for c in chunks)


def test_process_pdf_document():
    token = get_auth_token("pdf_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Create a minimal valid PDF in memory
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf_buffer = io.BytesIO()
    writer.write(pdf_buffer)
    pdf_buffer.seek(0)

    files = {"file": ("sample_paper.pdf", pdf_buffer, "application/pdf")}
    res_upload = client.post("/api/v1/documents/upload", headers=headers, files=files)
    assert res_upload.status_code == 201
    doc_id = res_upload.json()["id"]

    # Trigger processing on a blank page -> fails gracefully with empty text
    res_process = client.post(f"/api/v1/documents/{doc_id}/process", headers=headers)
    assert res_process.status_code == 200
    data = res_process.json()
    # It catches "No readable text" and marks as FAILED safely without crashing
    assert data["processing_status"] == "FAILED"
    assert "No readable text" in data["error_message"]


def test_process_nonexistent_or_unauthorized_document():
    token = get_auth_token("other_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    res = client.post("/api/v1/documents/invalid-uuid-1234/process", headers=headers)
    assert res.status_code == 404
